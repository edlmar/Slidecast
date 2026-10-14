"""Parse a .pptx into slide records (order, title, speaker notes, hidden flag)."""

from __future__ import annotations

import hashlib
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from slidecast.core.errors import DeckError


@dataclass(frozen=True)
class SlideRecord:
    index: int  # 0-based position in the deck
    title: str
    notes: str
    hidden: bool
    slide_id: int | None = None  # PowerPoint's id for the slide: stable across reordering

    @property
    def number(self) -> int:
        return self.index + 1


@dataclass
class DeckInfo:
    path: Path
    sha256: str
    visual_sha256: str = ""
    slides: list[SlideRecord] = field(default_factory=list)
    motion: DeckMotion = field(default_factory=lambda: DeckMotion())

    @property
    def name(self) -> str:
        return self.path.stem


def hash_file(path: Path, chunk: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            b = f.read(chunk)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


# Zip members that never appear in a slide image: speaker notes, comments, document
# properties, view state. Everything else - slide XML, layouts, masters, themes, media,
# slide order - is hashed so that editing narration text alone does not invalidate the
# render cache (a notes tweak must not cost a full LibreOffice re-render).
_NONVISUAL_PREFIXES = (
    "docProps/",
    "customXml/",
    "ppt/comments/",
    "ppt/handoutMasters/",
    "ppt/notesMasters/",
    "ppt/notesSlides/",
    "ppt/printerSettings/",
)
_NONVISUAL_FILES = {
    "ppt/commentAuthors.xml",
    "ppt/presProps.xml",
    "ppt/viewProps.xml",
}


def visual_hash(path: Path) -> str:
    """SHA-256 of the deck parts that decide what the slides look like.

    A .pptx rewritten with only different speaker notes (or just re-saved) keeps this
    hash, so notes edits and PowerPoint's own metadata churn do not invalidate cached
    slide images. Any change to slide content, layouts, masters, themes, media, or slide
    order changes it.
    """
    h = hashlib.sha256()
    try:
        with zipfile.ZipFile(path) as z:
            names = sorted(
                n
                for n in z.namelist()
                if not n.startswith(_NONVISUAL_PREFIXES) and n not in _NONVISUAL_FILES
            )
            for name in names:
                h.update(name.encode())
                h.update(z.read(name))
    except (OSError, zipfile.BadZipFile) as exc:
        raise DeckError(f"Could not hash {Path(path).name}: {exc}") from exc
    return h.hexdigest()


def _slide_title(slide: object, index: int) -> str:
    shapes = getattr(slide, "shapes", None)
    title_shape = getattr(shapes, "title", None) if shapes is not None else None
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        text = title_shape.text_frame.text.strip()
        if text:
            return " ".join(text.split())
    if shapes is not None:
        for shape in shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    first = text.splitlines()[0].strip()
                    return " ".join(first.split())[:120]
    return f"Slide {index + 1}"


def _slide_notes(slide: object) -> str:
    if not getattr(slide, "has_notes_slide", False):
        return ""
    notes_slide = slide.notes_slide  # type: ignore[attr-defined]
    tf = getattr(notes_slide, "notes_text_frame", None)
    if tf is None:
        return ""
    return str(tf.text or "")


def _slide_hidden(slide: object) -> bool:
    element = getattr(slide, "_element", None)
    if element is None:
        return False
    return bool(element.get("show") == "0")


VIDEO_SUFFIXES = (".mp4", ".mov", ".avi", ".wmv", ".m4v", ".mkv", ".webm", ".mpg", ".mpeg")
AUDIO_SUFFIXES = (".mp3", ".wav", ".m4a", ".wma", ".aac")
# An animated GIF is stored as a *picture*, not as media, and PowerPoint loops it forever.
ANIMATED_IMAGE_SUFFIXES = (".gif", ".webp", ".apng")


@dataclass(frozen=True)
class DeckMotion:
    """What a deck would lose if rendered as still images."""

    animated_slides: int = 0
    transition_slides: int = 0
    video_files: int = 0
    audio_files: int = 0
    animated_images: int = 0

    @property
    def any_motion(self) -> bool:
        return bool(
            self.animated_slides
            or self.transition_slides
            or self.video_files
            or self.animated_images
        )

    def describe(self) -> str:
        bits = []
        if self.animated_slides:
            bits.append(f"{self.animated_slides} slide(s) with animation")
        if self.transition_slides:
            bits.append(f"{self.transition_slides} with transitions")
        if self.video_files:
            bits.append(f"{self.video_files} embedded video(s)")
        if self.animated_images:
            bits.append(f"{self.animated_images} animated image(s)")
        if self.audio_files:
            bits.append(f"{self.audio_files} embedded audio file(s)")
        return ", ".join(bits) if bits else "no animations or media"


def inspect_motion(path: Path) -> DeckMotion:
    """Scan the .pptx for animation timing, slide transitions and embedded media.

    Reads the OOXML directly, so it is fast and needs neither PowerPoint nor LibreOffice.
    """
    import re

    animated = transitions = 0
    videos = audios = animated_images = 0
    try:
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                lower = name.lower()
                if lower.startswith("ppt/media/"):
                    if lower.endswith(VIDEO_SUFFIXES):
                        videos += 1
                    elif lower.endswith(AUDIO_SUFFIXES):
                        audios += 1
                    elif lower.endswith(ANIMATED_IMAGE_SUFFIXES) and _is_animated_image(
                        z.read(name)
                    ):
                        animated_images += 1
                elif re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
                    xml = z.read(name).decode("utf-8", "ignore")
                    if "<p:transition" in xml:
                        transitions += 1
                    # A bare <p:timing> node exists on many slides; require a real animation.
                    if re.search(r"<p:(anim|animEffect|animMotion|animRot|animScale)\b", xml):
                        animated += 1
    except (OSError, zipfile.BadZipFile) as exc:
        raise DeckError(f"Could not inspect {Path(path).name}: {exc}") from exc
    return DeckMotion(animated, transitions, videos, audios, animated_images)


def _is_animated_image(data: bytes) -> bool:
    """True for a multi-frame GIF/WebP. A single-frame one is just a picture."""
    import io

    try:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as img:
            return int(getattr(img, "n_frames", 1)) > 1
    except Exception:  # unreadable image: treat as static rather than guessing
        return False


def slides_with_media(path: Path) -> set[int]:
    """Slide numbers (1-based) that carry an audio or video shape.

    Such a slide has a natural length - PowerPoint plays its media to completion whatever the
    slide timing says - so its clip is rendered as-is. Every other slide is rendered to the
    length of its own narration instead.

    Matched on the relationship *type* rather than on filenames, so a video that lives outside
    the deck (a linked file, or a YouTube embed) counts too. Detection only chooses how long to
    render the slide; the rendered clip itself is what decides the layout, so a miss here costs
    a re-render rather than a wrong video.
    """
    import re

    kinds = re.compile(r'Type="[^"]*/relationships/(?:video|audio|media)"', re.IGNORECASE)
    out: set[int] = set()
    try:
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                m = re.fullmatch(r"ppt/slides/_rels/slide(\d+)\.xml\.rels", name)
                if not m:
                    continue
                if kinds.search(z.read(name).decode("utf-8", "ignore")):
                    out.add(int(m.group(1)))
    except (OSError, zipfile.BadZipFile) as exc:
        raise DeckError(f"Could not inspect {Path(path).name}: {exc}") from exc
    return out


def parse_deck(path: Path) -> DeckInfo:
    """Read slide order, titles, notes and hidden flags from a .pptx file."""
    path = Path(path)
    if not path.is_file():
        raise DeckError(f"Deck not found: {path}")
    if path.suffix.lower() != ".pptx":
        raise DeckError(f"Only .pptx decks are supported (got {path.suffix}). Save as .pptx first.")
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - dependency is mandatory
        raise DeckError("python-pptx is not installed") from exc

    try:
        prs = Presentation(str(path))
    except Exception as exc:  # python-pptx raises assorted exception types
        raise DeckError(f"Could not open {path.name}: {exc}") from exc

    slides: list[SlideRecord] = []
    for i, slide in enumerate(prs.slides):
        slides.append(
            SlideRecord(
                index=i,
                title=_slide_title(slide, i),
                notes=_slide_notes(slide),
                hidden=_slide_hidden(slide),
                slide_id=getattr(slide, "slide_id", None),
            )
        )
    return DeckInfo(
        path=path,
        sha256=hash_file(path),
        visual_sha256=visual_hash(path),
        slides=slides,
        motion=inspect_motion(path),
    )
